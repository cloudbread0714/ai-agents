"""A PPTX/DOCX → PDF renderer for tests that needs no LibreOffice.

DOCX: every paragraph in document order on A4 pages, wrapped at the text width, a new page
when the bottom margin is reached (so a longer resume adds pages). PPTX: one page per slide
in presentation order. Every text frame is drawn at its box with the
default text insets; lines wrap at the box width (character-greedy, like Korean text),
and text that does not fit keeps going past the box bottom, so overflow detection can be
tested deterministically.
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Emu

EMU_PER_PT = 12700
INSET_X, INSET_Y = 7.2, 3.6
DEFAULT_SIZE = 18.0

_FONT: pymupdf.Font | None = None


def _font() -> pymupdf.Font:
    global _FONT
    if _FONT is None:
        _FONT = pymupdf.Font("cjk")
    return _FONT


def _pt(v: int | Emu | None) -> float:
    return (v or 0) / EMU_PER_PT


def _text_shapes(shapes):
    for shape in shapes:
        if shape.shape_type is not None and shape.shape_type == 6:  # GROUP
            yield from _text_shapes(shape.shapes)
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape


def _wrap(text: str, width: float, size: float) -> list[str]:
    font, lines, line = _font(), [], ""
    for ch in text:
        if line and font.text_length(line + ch, fontsize=size) > width:
            lines.append(line)
            line = "" if ch == " " else ch
        else:
            line += ch
    return [*lines, line] if line else lines or [""]


A4_W, A4_H, MARGIN, BODY_SIZE = 595.0, 842.0, 50.0, 11.0


def render_docx(docx: Path, out_dir: Path) -> Path:
    from job_apply import docx_doc

    doc = pymupdf.open()
    page = writer = None
    y = A4_H
    for p, _where in docx_doc.walk(docx_doc.read_document(docx)):
        for line in _wrap(docx_doc.paragraph_text(p), A4_W - 2 * MARGIN, BODY_SIZE):
            if y + BODY_SIZE * 1.4 > A4_H - MARGIN:
                if writer is not None:
                    writer.write_text(page)
                page = doc.new_page(width=A4_W, height=A4_H)
                writer, y = pymupdf.TextWriter(page.rect), MARGIN
            y += BODY_SIZE * 1.4
            if line.strip():
                writer.append((MARGIN, y), line, font=_font(), fontsize=BODY_SIZE)
    if writer is not None:
        writer.write_text(page)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{docx.stem}.pdf"
    doc.save(out)
    doc.close()
    return out


def render(pptx: Path, out_dir: Path) -> Path:
    if pptx.suffix == ".docx":
        return render_docx(pptx, out_dir)
    prs = Presentation(str(pptx))
    width, height = _pt(prs.slide_width), _pt(prs.slide_height)
    doc = pymupdf.open()
    for slide in prs.slides:
        page = doc.new_page(width=width, height=height)
        writer = pymupdf.TextWriter(page.rect)
        for shape in _text_shapes(slide.shapes):
            left, top = _pt(shape.left), _pt(shape.top)
            avail = max(_pt(shape.width) - 2 * INSET_X, 10)
            y = top + INSET_Y
            for p in shape.text_frame.paragraphs:
                size = next((r.font.size.pt for r in p.runs if r.font.size), DEFAULT_SIZE)
                for line in _wrap(p.text, avail, size):
                    y += size * 1.2
                    if line.strip():
                        writer.append((left + INSET_X, y), line, font=_font(), fontsize=size)
        writer.write_text(page)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{pptx.stem}.pdf"
    doc.save(out)
    doc.close()
    return out


class FakeRenderer:
    """Renderer protocol stand-in. Records every call."""

    def __init__(self) -> None:
        self.calls: list[Path] = []

    def to_pdf(self, pptx: Path, out_dir: Path) -> Path:
        self.calls.append(pptx)
        return render(pptx, out_dir)
