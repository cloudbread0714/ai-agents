"""Sample originals and postings for tests. Everything here is fictional: the person
"홍길동", the companies and the numbers are made up (constitution: job-apply)."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pymupdf
from pptx import Presentation
from pptx.util import Pt

RESUME_LINES = [
    "홍길동",
    "백엔드 개발자 · hong@example.com",
    "경력",
    "가상테크 (2021-2024) 백엔드 개발",
    "Python FastAPI 기반 API 서버 설계와 운영",
    "PostgreSQL 쿼리 튜닝으로 응답 시간 30% 개선",
    "Docker 배포 파이프라인 구축",
    "기술",
    "Python, FastAPI, PostgreSQL, Docker, AWS",
]


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DOCX_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
    'officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-'
    'officedocument.wordprocessingml.styles+xml"/></Types>'
)
_DOCX_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/officeDocument" Target="word/document.xml"/></Relationships>'
)
_DOC_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/styles" Target="styles.xml"/></Relationships>'
)
_STYLES = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles xmlns:w="{W_NS}">'
    '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Pretendard" w:eastAsia="Pretendard"/>'
    '<w:sz w:val="22"/></w:rPr></w:rPrDefault></w:docDefaults></w:styles>'
)
BOLD = "<w:rPr><w:b/></w:rPr>"


def _p(*runs: tuple[str, bool]) -> str:
    body = "".join(
        f'<w:r>{BOLD if bold else ""}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
        for text, bold in runs
    )
    return f'<w:p><w:pPr><w:spacing w:after="60"/></w:pPr>{body}</w:p>'


def _tc(*paragraphs: str) -> str:
    return (
        '<w:tc><w:tcPr><w:tcW w:w="9000" w:type="dxa"/></w:tcPr>' + "".join(paragraphs) + "</w:tc>"
    )


def make_resume_docx(path: Path, lines: list[str] | None = None) -> Path:
    """A one-table resume: name and title, a 경력 heading, a table with one row per
    경력 line, then 기술. ``resume.txt`` lines match ``RESUME_LINES`` in order; extra
    ``lines`` become paragraphs at the end.

    Paragraph numbers: 0 name, 1 title, 2 "경력", 3–6 table rows 0–3, 7 "기술", 8 skills.
    """
    lines = lines or RESUME_LINES
    head, table_lines, tail = lines[:3], lines[3:7], lines[7:]
    rows = "".join(f"<w:tr>{_tc(_p((text, i == 0)))}</w:tr>" for i, text in enumerate(table_lines))
    table = (
        '<w:tbl><w:tblPr><w:tblW w:w="9000" w:type="dxa"/><w:tblBorders>'
        '<w:top w:val="single" w:sz="4"/></w:tblBorders></w:tblPr>'
        f'<w:tblGrid><w:gridCol w:w="9000"/></w:tblGrid>{rows}</w:tbl>'
    )
    body = (
        _p((head[0], True))
        + _p((head[1], False))
        + _p((head[2], True))
        + table
        + _p((tail[0], True))
        + "".join(_p((t, False)) for t in tail[1:])
        + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'
    )
    document = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}"><w:body>{body}</w:body></w:document>'
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _DOCX_TYPES)
        z.writestr("_rels/.rels", _DOCX_RELS)
        z.writestr("word/_rels/document.xml.rels", _DOC_RELS)
        z.writestr("word/styles.xml", _STYLES)
        z.writestr("word/document.xml", document)
    return path


def _png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    pix.clear_with(200)
    return pix.tobytes("png")


def _textbox(slide, name, left, top, width, height, paragraphs, size=18):
    """``paragraphs``: list of (runs, level) where runs is [(text, bold), ...]."""
    box = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(width), Pt(height))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    for i, (runs, level) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        for text, bold in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
    return box


def make_portfolio_pptx(path: Path) -> Path:
    """Five slides (720×540pt):

    1. title layout: title + subtitle
    2. "Intro" (2 runs: bold + normal), "Bullets" (3 paragraphs, levels 0/1/0), "Small" (tiny box)
    3. a picture and a "Caption" text box
    4. "Overlap A" / "Overlap B" text boxes that already overlap in the original
    5. a table and a group holding a "Grouped" text box
    """
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "홍길동 포트폴리오"
    title.placeholders[1].text = "백엔드 개발자"

    blank = prs.slide_layouts[6]
    s2 = prs.slides.add_slide(blank)
    _textbox(s2, "Intro", 40, 40, 640, 60, [([("핵심 역량: ", True), ("Python 백엔드", False)], 0)])
    _textbox(
        s2,
        "Bullets",
        40,
        120,
        640,
        200,
        [
            ([("API 서버 설계", False)], 0),
            ([("응답 시간 30% 개선", False)], 1),
            ([("PostgreSQL 운영", False)], 0),
        ],
    )
    _textbox(s2, "Small", 40, 360, 200, 30, [([("짧은 메모", False)], 0)], size=14)

    s3 = prs.slides.add_slide(blank)
    import io

    s3.shapes.add_picture(io.BytesIO(_png()), Pt(40), Pt(40), Pt(200), Pt(200))
    _textbox(s3, "Caption", 260, 40, 400, 60, [([("아키텍처 다이어그램", False)], 0)])

    s4 = prs.slides.add_slide(blank)
    _textbox(s4, "Overlap A", 40, 40, 300, 60, [([("장식 제목 A", False)], 0)])
    _textbox(s4, "Overlap B", 120, 50, 300, 60, [([("장식 제목 B", False)], 0)])

    s5 = prs.slides.add_slide(blank)
    table = s5.shapes.add_table(2, 2, Pt(40), Pt(40), Pt(300), Pt(80)).table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    group = s5.shapes.add_group_shape()
    group.name = "Group"
    inner = group.shapes.add_textbox(Pt(40), Pt(200), Pt(300), Pt(60))
    inner.name = "Grouped"
    inner.text_frame.text = "그룹 안 텍스트"

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


_LONG = (
    "가상테크는 백엔드 개발자를 찾습니다. 주요 업무는 Python 기반 API 서버 개발과 운영입니다. "
    "자격 요건: Python 3년 이상, PostgreSQL 경험, Docker 사용 경험. "
    "우대 사항: Kubernetes 운영 경험, AWS 경험, 대용량 트래픽 처리 경험. "
    "근무지: 서울, 정규직. 전형: 서류 → 과제 → 면접. "
)


def posting_text(repeat: int = 4) -> str:
    return _LONG * repeat


def make_posting_html(kind: str) -> str:
    if kind == "jsonld":
        data = {
            "@context": "https://schema.org",
            "@type": "JobPosting",
            "title": "백엔드 개발자",
            "hiringOrganization": {"@type": "Organization", "name": "가상테크"},
            "validThrough": "2026-11-10",
            "description": f"<p>{posting_text()}</p>",
        }
        return (
            "<html><head><script type='application/ld+json'>"
            + json.dumps(data, ensure_ascii=False)
            + "</script></head><body><nav>메뉴 메뉴 메뉴</nav></body></html>"
        )
    if kind == "next_data":
        data = {
            "props": {"pageProps": {"job": {"title": "데이터 엔지니어", "body": posting_text()}}}
        }
        return (
            "<html><body><div id='__next'></div><script id='__NEXT_DATA__' "
            "type='application/json'>"
            + json.dumps(data, ensure_ascii=False)
            + "</script></body></html>"
        )
    if kind == "plain":
        paragraphs = "".join(f"<p>{i}번째 안내. {posting_text(1)}</p>" for i in range(1, 5))
        return f"<html><body><article><h1>백엔드 개발자</h1>{paragraphs}</article></body></html>"
    if kind == "login_wall":
        return (
            "<html><body><main><h1>로그인</h1>"
            "<p>로그인이 필요합니다. 로그인 후 공고를 확인하세요.</p>"
            "<form><input name='id'><input name='pw'></form></main></body></html>"
        )
    if kind == "too_short":
        body = "<p>백엔드 개발자 모집. 자세한 내용은 문의.</p>"
        return f"<html><body><article>{body}</article></body></html>"
    raise ValueError(kind)
