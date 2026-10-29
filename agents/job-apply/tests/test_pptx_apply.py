import io

import pytest
from lxml import etree
from pptx import Presentation

from ja_support import session
from ja_support.samples import make_portfolio_pptx
from job_apply.errors import PlanError
from job_apply.models import parse_plan
from job_apply.pptx_apply import apply, edited_slides
from job_apply.pptx_dump import dump

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _edit(slide=2, shape_id=2, paragraph=0, **change):
    return {
        "slide": slide,
        "shape_id": shape_id,
        "paragraph": paragraph,
        "source": f"portfolio.json:s{slide}/sh{shape_id}/p{paragraph}",
        "reason": "테스트용 변경 사유",
        **change,
    }


def _run(tmp_path, edits, order=None):
    deck = make_portfolio_pptx(tmp_path / "deck.pptx")
    plan = parse_plan(session.plan(edits=edits, order=order), 5)
    prs = Presentation(io.BytesIO(apply(deck, plan)))
    return prs, dump(prs)


def _shape(prs, slide, shape_id):
    return next(s for s in prs.slides[slide - 1].shapes if s.shape_id == shape_id)


def test_text_goes_into_first_run_keeping_its_format(tmp_path):
    prs, data = _run(tmp_path, [_edit(text="핵심 역량: API 서버")])
    p = _shape(prs, 2, 2).text_frame.paragraphs[0]
    assert [r.text for r in p.runs] == ["핵심 역량: API 서버"]
    assert p.runs[0].font.bold is True  # first run's a:rPr kept


def test_runs_must_match_run_count(tmp_path):
    prs, _ = _run(tmp_path, [_edit(runs=["역량: ", "API"])])
    p = _shape(prs, 2, 2).text_frame.paragraphs[0]
    assert [(r.text, r.font.bold) for r in p.runs] == [("역량: ", True), ("API", False)]
    with pytest.raises(PlanError, match="run 개수"):
        _run(tmp_path, [_edit(runs=["하나뿐"])])


def test_insert_after_clones_format_and_delete(tmp_path):
    prs, data = _run(
        tmp_path,
        [
            _edit(shape_id=3, paragraph=1, op={"kind": "insert_after", "text": "새 항목"}),
            _edit(shape_id=3, paragraph=1, op={"kind": "insert_after", "text": "그다음 항목"}),
            _edit(shape_id=3, paragraph=2, op={"kind": "delete"}),
        ],
    )
    paragraphs = _shape(prs, 2, 3).text_frame.paragraphs
    assert [(p.text, p.level) for p in paragraphs] == [
        ("API 서버 설계", 0),
        ("응답 시간 30% 개선", 1),
        ("새 항목", 1),
        ("그다음 항목", 1),
    ]
    clone = etree.tostring(paragraphs[2].runs[0]._r.find(f"{A}rPr"))
    assert clone == etree.tostring(paragraphs[1].runs[0]._r.find(f"{A}rPr"))


def test_cannot_delete_last_paragraph(tmp_path):
    with pytest.raises(PlanError, match="마지막 문단"):
        _run(tmp_path, [_edit(shape_id=4, op={"kind": "delete"})])


def test_slide_order_only_reorders(tmp_path):
    prs, data = _run(tmp_path, [_edit(text="x")], order=[1, 3, 2, 4, 5])
    assert [s["shapes"][0]["name"] for s in data["slides"][1:3]] == ["Caption", "Intro"]
    assert edited_slides(
        parse_plan(session.plan(edits=[_edit(text="x")], order=[1, 3, 2, 4, 5]), 5)
    ) == [3]


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        (_edit(shape_id=99, text="x"), "shape_id 99"),
        (_edit(paragraph=5, text="x"), "문단 5"),
        (_edit(slide=5, shape_id=2, text="x"), "표"),
        (_edit(slide=3, shape_id=2, text="x"), "텍스트가 없어요"),
    ],
)
def test_unknown_targets_are_plan_errors(tmp_path, edit, message):
    with pytest.raises(PlanError, match=message):
        _run(tmp_path, [edit])
    assert PlanError.exit_code == 2


def test_group_text_is_editable(tmp_path):
    prs, data = _run(tmp_path, [_edit(slide=5, shape_id=4, text="그룹 문구 변경")])
    assert data["slides"][4]["shapes"][1]["paragraphs"][0]["text"] == "그룹 문구 변경"
