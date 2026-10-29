"""Apply ``tailor.json`` portfolio edits to the base copy (research R3).

Only paragraph text and slide order change. New text goes into the paragraph's first run,
keeping its ``a:rPr``; inserted paragraphs are clones of a sibling in the same shape. The
design guard (design_guard.py) checks the result before anything is written.
"""

from __future__ import annotations

import copy
import io
from pathlib import Path

from pptx import Presentation

from job_apply.errors import PlanError
from job_apply.models import PortfolioEdit, TailorPlan
from job_apply.pptx_dump import walk

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_RUN_LIKE = (f"{A}r", f"{A}br", f"{A}fld")


def _find_shape(slide, shape_id: int):
    for shape, _box in walk(slide.shapes):
        if shape.shape_id == shape_id:
            return shape
    return None


def _target(prs, edit: PortfolioEdit, where: str):
    slides = list(prs.slides)
    if not 1 <= edit.slide <= len(slides):
        raise PlanError(f"{where}: 슬라이드 {edit.slide}이(가) 없어요")
    shape = _find_shape(slides[edit.slide - 1], edit.shape_id)
    if shape is None:
        raise PlanError(f"{where}: 슬라이드 {edit.slide}에 shape_id {edit.shape_id}가 없어요")
    if getattr(shape, "has_table", False) and shape.has_table:
        raise PlanError(f"{where}: 표는 이번 버전에서 고칠 수 없어요")
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        raise PlanError(f"{where}: shape_id {edit.shape_id}에는 텍스트가 없어요")
    paragraphs = shape.text_frame.paragraphs
    if edit.paragraph >= len(paragraphs):
        raise PlanError(
            f"{where}: shape_id {edit.shape_id}에 문단 {edit.paragraph}이(가) 없어요 "
            f"(문단 {len(paragraphs)}개)"
        )
    return shape, paragraphs[edit.paragraph]._p


def _set_text(p, text: str) -> None:
    """Text into the first run (its a:rPr kept); every other run, break and field removed."""
    runs = [el for el in p if el.tag in _RUN_LIKE]
    first = next((el for el in runs if el.tag == f"{A}r"), None)
    for el in runs:
        if el is not first:
            p.remove(el)
    if first is None:
        first = p.makeelement(f"{A}r", {})
        end = p.find(f"{A}endParaRPr")
        if end is not None:
            end.addprevious(first)
        else:
            p.append(first)
    t = first.find(f"{A}t")
    if t is None:
        t = first.makeelement(f"{A}t", {})
        first.append(t)
    t.text = text


def _set_runs(p, texts: list[str], where: str) -> None:
    runs = p.findall(f"{A}r")
    if len(runs) != len(texts):
        raise PlanError(f"{where}: runs는 원본 run 개수({len(runs)})와 같아야 해요")
    for r, text in zip(runs, texts, strict=True):
        r.find(f"{A}t").text = text


def apply(base_pptx: Path, plan: TailorPlan) -> bytes:
    prs = Presentation(str(base_pptx))
    resolved = [
        (edit, *_target(prs, edit, f"edits[{i}]")) for i, edit in enumerate(plan.portfolio.edits)
    ]
    # Text first, then inserts (after the original paragraph, in plan order), deletes last,
    # so every edit addresses the base paragraph numbering.
    for i, (edit, _shape, p) in enumerate(resolved):
        if edit.text is not None:
            _set_text(p, edit.text)
        elif edit.runs is not None:
            _set_runs(p, edit.runs, f"edits[{i}]")
    last_inserted: dict[int, object] = {}
    for edit, _shape, p in resolved:
        if edit.op and edit.op.kind == "insert_after":
            clone = copy.deepcopy(p)
            _set_text(clone, edit.op.text or "")
            anchor = last_inserted.get(id(p), p)
            anchor.addnext(clone)
            last_inserted[id(p)] = clone
    for i, (edit, _shape, p) in enumerate(resolved):
        if edit.op and edit.op.kind == "delete":
            body = p.getparent()
            if len(body.findall(f"{A}p")) <= 1:
                raise PlanError(f"edits[{i}]: 도형의 마지막 문단은 지울 수 없어요")
            body.remove(p)

    ids = list(prs.slides._sldIdLst)
    for el in ids:
        prs.slides._sldIdLst.remove(el)
    for n in plan.portfolio.slide_order:
        prs.slides._sldIdLst.append(ids[n - 1])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def edited_slides(plan: TailorPlan) -> list[int]:
    """Result-order slide numbers whose text changed (for previews)."""
    base = {e.slide for e in plan.portfolio.edits}
    order = plan.portfolio.slide_order
    return sorted(i for i, n in enumerate(order, start=1) if n in base)
