"""``portfolio.json``: the text structure of a deck, the only PPTX view the session gets
(contracts/cli.md init). Shape ids are python-pptx ``shape_id`` (``p:cNvPr/@id``), paragraph
indexes are 0-based within the shape's text frame."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_PT = 12700


def _pt(v: float) -> float:
    return round(v / EMU_PER_PT, 1)


def _xfrm(shape) -> tuple[float, float, float, float]:
    return (
        float(shape.left or 0),
        float(shape.top or 0),
        float(shape.width or 0),
        float(shape.height or 0),
    )


def _group_map(group, outer):
    """Map a child box in the group's child space to slide space (EMU)."""
    xfrm = group._element.grpSpPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm"
    )
    gx, gy, gw, gh = _xfrm(group)
    if xfrm is None or xfrm.find("{*}chOff") is None:
        cx, cy, cw, ch = gx, gy, gw, gh
    else:
        off, ext = xfrm.find("{*}chOff"), xfrm.find("{*}chExt")
        cx, cy = float(off.get("x")), float(off.get("y"))
        cw, ch = float(ext.get("cx")) or gw, float(ext.get("cy")) or gh
    sx, sy = (gw / cw if cw else 1.0), (gh / ch if ch else 1.0)

    def to_slide(box):
        x, y, w, h = box
        return outer((gx + (x - cx) * sx, gy + (y - cy) * sy, w * sx, h * sy))

    return to_slide


def walk(shapes, outer=lambda b: b):
    """Yield (shape, slide-space box in EMU) for every shape, entering groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes, _group_map(shape, outer))
        else:
            yield shape, outer(_xfrm(shape))


def box_pt(box) -> list[float]:
    x, y, w, h = box
    return [_pt(x), _pt(y), _pt(x + w), _pt(y + h)]


def dump(prs) -> dict[str, Any]:
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else None
        shapes = []
        for shape, box in walk(slide.shapes):
            entry: dict[str, Any] = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "box_pt": box_pt(box),
            }
            if getattr(shape, "has_table", False) and shape.has_table:
                entry["table"] = True
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                entry["paragraphs"] = [
                    {
                        "paragraph": i,
                        "text": p.text,
                        "runs": [r.text for r in p.runs],
                        "level": p.level,
                    }
                    for i, p in enumerate(shape.text_frame.paragraphs)
                ]
            else:
                continue
            shapes.append(entry)
        slides.append({"slide": number, "title": title, "shapes": shapes})
    return {"slide_count": len(slides), "slides": slides}


def dump_file(pptx: Path, out: Path | None = None) -> dict[str, Any]:
    data = dump(Presentation(str(pptx)))
    if out is not None:
        out.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return data


def all_text(data: dict[str, Any]) -> list[str]:
    return [
        p["text"]
        for s in data["slides"]
        for sh in s["shapes"]
        for p in sh.get("paragraphs", [])
        if p["text"]
    ]
